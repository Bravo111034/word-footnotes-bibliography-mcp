import zipfile
from lxml import etree

import docx
import pytest

from docx_fnbib import api
from docx_fnbib.footnotes import FootnoteManager

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


@pytest.fixture
def sample(tmp_path):
    p = tmp_path / "doc.docx"
    api.create_document(str(p), title="Test", paragraphs=[
        "First body paragraph with a claim.",
        "Second paragraph mentioning Napoleon in passing.",
    ])
    return str(p)


def test_add_footnote_creates_part_and_reference(sample):
    res = api.add_footnote(sample, 1, "See Smith 2020, p. 5.")
    assert res["footnote_id"] == 1

    with zipfile.ZipFile(sample) as z:
        names = z.namelist()
        assert "word/footnotes.xml" in names
        fx = etree.fromstring(z.read("word/footnotes.xml"))
        dx = etree.fromstring(z.read("word/document.xml"))
        ct = z.read("[Content_Types].xml").decode()

    real = [f for f in fx.findall(f"{{{W}}}footnote") if not f.get(f"{{{W}}}type")]
    assert len(real) == 1
    assert real[0].get(f"{{{W}}}id") == "1"
    assert "See Smith 2020" in "".join(fx.itertext())

    refs = dx.findall(f".//{{{W}}}footnoteReference")
    assert len(refs) == 1 and refs[0].get(f"{{{W}}}id") == "1"
    assert "footnotes+xml" in ct


def test_multiple_footnotes_increment(sample):
    api.add_footnote(sample, 0, "One.")
    api.add_footnote(sample, 1, "Two.")
    api.add_footnote(sample, 1, "Three.")
    notes = api.list_footnotes(sample)
    assert [n["id"] for n in notes] == [1, 2, 3]
    assert notes[2]["text"] == "Three."


def test_anchor_places_reference_after_word(sample):
    # paragraph 0 is the title heading; "Napoleon" is in paragraph 2
    api.add_footnote(sample, 2, "Emperor of the French.", anchor="Napoleon")
    doc = docx.Document(sample)
    para = doc.paragraphs[2]
    # the run carrying the reference should sit right after the text 'Napoleon'
    xml = para._p.xml
    idx_napoleon = xml.find("Napoleon")
    idx_ref = xml.find("footnoteReference")
    assert idx_napoleon < idx_ref
    assert api.list_footnotes(sample)[0]["text"] == "Emperor of the French."


def test_reopen_and_append_keeps_existing(sample):
    api.add_footnote(sample, 0, "First.")
    # reopen through a fresh manager
    doc = docx.Document(sample)
    fm = FootnoteManager(doc)
    assert fm.next_id() == 2
    fm.insert(doc.paragraphs[1], "Second.")
    doc.save(sample)
    assert len(api.list_footnotes(sample)) == 2
