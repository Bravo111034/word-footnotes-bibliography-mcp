"""Insert citations and build a bibliography into a document.

Two rendering modes:

* ``"field"`` – write Word ``CITATION`` / ``BIBLIOGRAPHY`` fields backed by the
  customXml source store. Word recomputes them and lets the user switch styles.
* ``"text"`` – render formatted citation/reference text ourselves (portable,
  no Word recalculation needed). Uses :mod:`.formatting`.
"""
from __future__ import annotations

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.text.paragraph import Paragraph
from lxml import etree

from .ooxml import qn
from . import formatting
from .footnotes import FootnoteManager
from .sources import Source, load_sources, save_sources
from .styles import ensure_styles

_WORD_STYLE_MAP = {
    "apa": "APA",
    "mla": "MLASeventhEdition",
    "chicago-author-date": "Chicago",
    "chicago-notes": "Chicago",
}


# --------------------------------------------------------------------------
# field helpers
# --------------------------------------------------------------------------
def _add_field(paragraph: Paragraph, instr: str, cached: str = "") -> None:
    p = paragraph._p

    def run(*children):
        r = etree.SubElement(p, qn("w:r"))
        for c in children:
            r.append(c)
        return r

    fld_begin = etree.Element(qn("w:fldChar"))
    fld_begin.set(qn("w:fldCharType"), "begin")
    run(fld_begin)

    it = etree.Element(qn("w:instrText"))
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    run(it)

    fld_sep = etree.Element(qn("w:fldChar"))
    fld_sep.set(qn("w:fldCharType"), "separate")
    run(fld_sep)

    if cached:
        t = etree.Element(qn("w:t"))
        t.set(qn("xml:space"), "preserve")
        t.text = cached
        run(t)

    fld_end = etree.Element(qn("w:fldChar"))
    fld_end.set(qn("w:fldCharType"), "end")
    run(fld_end)


def _add_segments(paragraph: Paragraph, segs, base_style: str | None = None) -> None:
    for text, italic in segs:
        r = paragraph.add_run(text)
        if italic:
            r.italic = True


# --------------------------------------------------------------------------
# public API
# --------------------------------------------------------------------------
def add_citation(
    document,
    paragraph: Paragraph,
    tag: str,
    *,
    style: str = "apa",
    mode: str = "text",
    page: str | None = None,
    prefix: str | None = None,
    lcid: int = 1033,
) -> str:
    """Append a citation for ``tag`` to ``paragraph``. Returns the plain text."""
    sources = {s.tag: s for s in load_sources(document)}
    if tag not in sources:
        raise KeyError(f"no source with tag {tag!r}; add it first")
    src = sources[tag]

    if style == "chicago-notes":
        # footnote-style citation
        fm = FootnoteManager(document)
        note_text = formatting.plain(
            formatting.in_text_citation(src, style, page=page)
        )
        fm.insert(paragraph, (prefix + " " if prefix else "") + note_text)
        return note_text

    if mode == "field":
        instr = f' CITATION {tag} \\l {lcid} '
        if page:
            instr += f'\\p {page} '
        if prefix:
            instr += f'\\f "{prefix} " '
        cached = formatting.plain(formatting.in_text_citation(src, style, page=page))
        _add_field(paragraph, instr, cached)
        return cached

    segs = formatting.in_text_citation(src, style, page=page)
    if prefix:
        segs = [(prefix + " ", False)] + list(segs)
    _add_segments(paragraph, segs)
    return formatting.plain(segs)


def add_bibliography(
    document,
    *,
    style: str = "apa",
    mode: str = "text",
    heading: str | None = "References",
    heading_level: int = 1,
    lcid: int = 1033,
) -> int:
    """Append a bibliography/works-cited section. Returns entry count."""
    ensure_styles(document, ["Bibliography"])
    sources = load_sources(document)

    if heading:
        document.add_heading(heading, level=heading_level)

    if mode == "field":
        p = document.add_paragraph()
        instr = f' BIBLIOGRAPHY \\l {lcid} '
        _add_field(p, instr, "")
        # persist the selected style name so Word renders it
        _set_source_style(document, sources, style)
        return len(sources)

    ordered = sorted(sources, key=formatting.sort_key)
    for src in ordered:
        p = document.add_paragraph(style="Bibliography")
        segs = formatting.bibliography_entry(src, style)
        _add_segments(p, segs)
    return len(ordered)


def _set_source_style(document, sources: list[Source], style: str) -> None:
    from .sources import _find_sources_part  # noqa

    part, root = _find_sources_part(document)
    if root is None:
        save_sources(document, sources)
        part, root = _find_sources_part(document)
    if root is not None:
        root.set("SelectedStyle", _WORD_STYLE_MAP.get(style, "APA"))
        root.set("StyleName", _WORD_STYLE_MAP.get(style, "APA"))
        part._blob = etree.tostring(
            root, xml_declaration=True, encoding="UTF-8", standalone=True
        )
